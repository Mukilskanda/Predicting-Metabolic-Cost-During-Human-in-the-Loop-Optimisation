"""Build docs/MetaCost_Presentation.pptx.   python docs/build_slides.py"""
from pathlib import Path

import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "results" / "figures"
TAB = ROOT / "results" / "tables"
OUT = ROOT / "docs" / "MetaCost_Presentation.pptx"

NAVY = RGBColor(0x0F, 0x1B, 0x2D)
NAVY2 = RGBColor(0x1C, 0x2C, 0x45)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
INK = RGBColor(0x14, 0x17, 0x1C)
INK2 = RGBColor(0x55, 0x5B, 0x66)
TINT = RGBColor(0xF1, 0xF5, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xBF, 0xD4, 0xF2)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
_n = [0]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """runs: str, or list of paragraphs; a paragraph is a str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.name = o.get("font", font)
            r.font.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        parts = it if isinstance(it, list) else [(it, {})]
        r = p.add_run()
        r.text = "▸  "
        r.font.size, r.font.color.rgb, r.font.name, r.font.bold = Pt(size), BLUE, BODY, True
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.name = BODY
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def image(slide, name, x, y, w=None, h=None):
    """Place an image fitting inside w x h (inches), centred in that box."""
    p = FIG / name
    iw, ih = PILImage.open(p).size
    r = iw / ih
    if w and h:
        if w / h > r:
            ww, hh = h * r, h
        else:
            ww, hh = w, w / r
        x, y = x + (w - ww) / 2, y + (h - hh) / 2
    elif w:
        ww, hh = w, w / r
    else:
        ww, hh = h * r, h
    return slide.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(ww), Inches(hh))


def chip(slide, x, y, n):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.5), Inches(0.5))
    c.fill.solid()
    c.fill.fore_color.rgb = BLUE
    c.line.fill.background()
    tf = c.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = str(n)
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(16), True, WHITE, BODY


def content_slide(title, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    _n[0] += 1
    if kicker:
        text(s, 0.6, 0.38, 10, 0.3, kicker.upper(), size=12, color=BLUE, bold=True)
    text(s, 0.6, 0.68, 12.1, 0.8, title, size=30, color=NAVY, font=HEAD, bold=True)
    text(s, 0.6, 7.02, 8, 0.3, "MetaCost · Predicting metabolic cost during HILO", size=10, color=INK2)
    text(s, 12.0, 7.02, 0.73, 0.3, str(_n[0] + 1), size=10, color=INK2, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def stat(slide, x, y, w, big, label, color=BLUE, big_size=40, fill=TINT, label_color=INK2):
    box(slide, x, y, w, 1.55, fill=fill)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.75, big, size=big_size, color=color, font=HEAD, bold=True)
    text(slide, x + 0.25, y + 0.95, w - 0.5, 0.55, label, size=13, color=label_color)


# ---------------------------------------------------------------- numbers
e3 = pd.read_csv(TAB / "E3_model_zoo.csv")
e1 = pd.read_csv(TAB / "E1_reproduction.csv")
M = lambda s, m, p="random", fs="All features": e3[(e3.subject == s) & (e3.model == m) & (e3.protocol == p)
                                                   & (e3.feature_set == fs)].iloc[0]

# ================================================================ 1. title
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
box(s, 0.6, 1.25, 1.1, 0.42, fill=BLUE, radius=0.5)
text(s, 0.6, 1.25, 1.1, 0.42, "ML · #11", size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 1.95, 11.5, 2.2, ["Predicting Metabolic Cost", "During Human-in-the-Loop Optimisation"],
     size=44, color=WHITE, font=HEAD, bold=True)
text(s, 0.6, 4.05, 11, 0.9, "Replacing a slow, lab-only face-mask measurement with wearable EMG and gait "
     "sensors, for exoskeletons that tune themselves outside the lab", size=18, color=ICE)
text(s, 0.6, 5.55, 8, 1.0, [[("Team:  ", {"bold": True, "color": WHITE}),
                             ("Mukil Skanda A (PES1UG24CS281)   ·   Likhith M (PES1UG24CS282)", {"color": ICE})],
                            [("UE24CS352A Machine Learning · Mini-Project Review · October 2026",
                              {"color": ICE, "size": 14})]], size=16, spacing=6)
s.notes_slide.notes_text_frame.text = (
    "Introduce the team. One line: we predict how much energy a person burns while walking in an exoskeleton, "
    "from sensors they could wear anywhere, instead of from a lab-only breathing mask.")

# ================================================================ 2. problem
s = content_slide("Exoskeleton tuning is bottlenecked by a breathing mask",
                  "The problem",
                  "Explain HILO: CMA-ES proposes a torque profile every 2 min and keeps those with the lowest metabolic "
                  "cost (Zhang et al., Science 2017). Metabolic cost comes from a mask (indirect calorimetry): slow, noisy, "
                  "lab-only. If we can predict it from wearable signals, HILO becomes faster and portable.")
# loop diagram
steps = [("CMA-ES optimiser", "proposes new torque profile"), ("Exoskeleton", "applies it for 2 min"),
         ("Human walks", "body adapts"), ("Metabolic cost", "measured via face mask")]
cx = [0.6, 3.55, 6.5, 9.45]
for i, ((t1, t2), x) in enumerate(zip(steps, cx)):
    hot = i == 3
    box(s, x, 1.9, 2.75, 1.35, fill=ORANGE if hot else TINT)
    text(s, x + 0.2, 2.05, 2.35, 0.45, t1, size=17, bold=True, color=WHITE if hot else NAVY)
    text(s, x + 0.2, 2.55, 2.35, 0.6, t2, size=14, color=WHITE if hot else INK2)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.8), Inches(2.42), Inches(0.2), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = INK2; a.line.fill.background()
text(s, 0.6, 3.38, 12, 0.4, "↻  repeated 36× per 72-min session, feedback = metabolic cost", size=14, color=INK2,
     italic=True)
stat(s, 0.6, 4.1, 3.85, "~1 min", "to settle after every change in control law, so it is slow feedback", ORANGE)
stat(s, 4.75, 4.1, 3.85, "≈ 5 %", "irreducible noise in the measurement (MSE floor ≈ 0.0025)", ORANGE)
stat(s, 8.9, 4.1, 3.85, "Lab only", "needs a face mask, a metabolic cart and fasting subjects", ORANGE)
text(s, 0.6, 6.0, 12.1, 0.8, [[("Our goal: ", {"bold": True, "color": NAVY}),
                               ("learn  f(EMG, gait, controls) → metabolic cost  so the mask can eventually be "
                                "replaced by sensors a person could wear anywhere.", {})]], size=17)

# ================================================================ 3. dataset
s = content_slide("Real lab data: 2 people, 5 days, 36 assistance conditions per day", "Dataset",
                  "Data from the Stanford Biomechatronics Lab HILO study, as released with the CS229 paper we were given. "
                  "Each row is the last 30 s of a 2-min condition. One Subject-2 row had a negative (impossible) metabolic "
                  "rate and is removed, as in the original. Point at the plot: cost drops ~40% across days because the user "
                  "learns to walk with the exo. The data are non-stationary, which matters later.")
stat(s, 0.6, 1.65, 2.0, "359", "samples\n(180 + 179)", big_size=34)
stat(s, 2.8, 1.65, 2.0, "29", "features per\nsample", big_size=34)
stat(s, 5.0, 1.65, 2.0, "0.0025", "noise-floor\nMSE", big_size=30)
feat = [("16  EMG", "8 muscles per leg; band-pass → rectify → RMS, normalised each day", BLUE),
        ("9  Step / gait", "peak ground force, ankle dorsi/plantar-flexion, step time, step width", AQUA),
        ("4  Control", "exo torque profile: peak torque, peak time, rise time, fall time", RGBColor(0xED, 0xA1, 0)),
        ("1  Target", "metabolic rate - standing baseline, ÷ normal walking (1.0 = no exo)", ORANGE)]
for i, (a, b, c) in enumerate(feat):
    y = 3.5 + i * 0.85
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(y + 0.08), Inches(0.28), Inches(0.28))
    d.fill.solid(); d.fill.fore_color.rgb = c; d.line.fill.background()
    text(s, 1.05, y, 5.6, 0.8, [[(a, {"bold": True, "color": NAVY, "size": 17})], [(b, {"size": 13, "color": INK2})]])
image(s, "F1_dataset_overview.png", 7.25, 1.55, w=5.55, h=5.3)

# ================================================================ 4. pipeline
s = content_slide("End-to-end pipeline", "Methodology",
                  "Walk through left to right. Pre-processing was done by the lab: EMG filtering and day normalisation. "
                  "Our contribution starts at the feature matrix: a leakage-free model zoo, two validation protocols and "
                  "analyses aimed at what the optimiser actually needs.")
image(s, "F13_pipeline.png", 0.6, 1.6, w=12.1)
cards = [("Leakage-free", "Scaling and every hyper-parameter search happen inside training folds (nested CV). "
          "The paper tuned λ on its test folds."),
         ("Two protocols", "Random 5-fold ×5 (the paper's, for comparison) and leave-one-day-out "
          "(predict a whole new session, the realistic case)."),
         ("Two layers", "A faithful 1:1 Python port of the paper's MATLAB scripts (baseline), plus the "
          "extended study: 12 experiments, 16 figures, 10 unit tests.")]
for i, (t1, t2) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 4.85, 3.85, 1.85)
    chip(s, x + 0.25, 5.05, i + 1)
    text(s, x + 0.9, 5.08, 2.8, 0.45, t1, size=17, bold=True, color=NAVY)
    text(s, x + 0.25, 5.65, 3.4, 1.0, t2, size=13, color=INK2)

# ================================================================ 5. models
s = content_slide("13 models, including the paper's network rebuilt from scratch", "Models",
                  "The paper used MATLAB trainbr. We re-implemented it in NumPy: Levenberg-Marquardt steps plus MacKay's "
                  "evidence updates for alpha, beta and gamma, so regularisation is learned from the training data with no "
                  "validation split. The Jacobian is closed form and unit-tested against finite differences. The neuron "
                  "sweep (right) reproduces the paper's Fig. 1: 1-2 neurons are best. The Gaussian Process adds calibrated "
                  "uncertainty.")
groups = [("Linear", "OLS · LASSO · Ridge · Elastic Net · PLS"), ("Kernel / instance", "SVR (RBF) · kNN · Gaussian Process"),
          ("Trees", "Random Forest · Extra Trees · Gradient Boosting"), ("Neural", "Bayesian-regularised NN (trainbr)"),
          ("Ensemble", "average of GP + SVR + LASSO + Extra Trees")]
for i, (a, b) in enumerate(groups):
    y = 1.65 + i * 0.66
    text(s, 0.6, y, 2.2, 0.5, a, size=15, bold=True, color=BLUE)
    text(s, 2.75, y, 4.0, 0.6, b, size=15, color=INK)
box(s, 0.6, 5.1, 6.1, 1.7, fill=NAVY)
text(s, 0.85, 5.22, 5.7, 1.5, [[("Bayesian-regularised NN, from scratch", {"bold": True, "color": WHITE, "size": 15})],
                               [("F(w) = β·E_D + α·E_W", {"color": ICE, "size": 14, "font": "Cambria"})],
                               [("γ = P − α·tr(H⁻¹),  α = γ / 2E_W,  β = (N − γ) / 2E_D",
                                 {"color": ICE, "size": 14, "font": "Cambria"})],
                               [("Levenberg-Marquardt · closed-form Jacobian · picks 1-2 neurons", {"color": WHITE, "size": 12})]],
     spacing=3)
image(s, "F3_neuron_sweep.png", 7.0, 1.55, w=5.75, h=5.3)

# ================================================================ 6. headline results
s = content_slide("We beat the original study on both subjects", "Results",
                  "Same protocol as the paper: random 5-fold CV, repeated 5 times and properly nested. S1: 0.0079 vs 0.0089 "
                  "(11% lower). S2: 0.0109 vs 0.0168 (35% lower). In plain terms, about 9-10% error relative to "
                  "normal-walking cost, versus 19-26% for predicting the mean. The noise floor is 0.0025, so headroom remains.")
stat(s, 0.6, 1.7, 2.95, "0.0079", "Subject 1 MSE (ours)\npaper: 0.0089", BLUE, big_size=38)
stat(s, 3.75, 1.7, 2.95, "−11 %", "error vs paper\non Subject 1", BLUE, big_size=38)
stat(s, 0.6, 3.45, 2.95, "0.0109", "Subject 2 MSE (ours)\npaper: 0.0168", ORANGE, big_size=38)
stat(s, 3.75, 3.45, 2.95, "−35 %", "error vs paper\non Subject 2", ORANGE, big_size=38)
stat(s, 0.6, 5.2, 6.1, "R² = 0.88  /  0.71", "variance explained (S1 / S2), out-of-fold", NAVY, big_size=30)
image(s, "F6_pred_vs_actual.png", 7.0, 1.55, w=5.75, h=5.3)

# ================================================================ 7. model zoo
s = content_slide("Random splits flatter models: unseen days are the real test", "Results · validation",
                  "Dark bars use the paper's protocol; light bars use leave-one-day-out. Every model gets worse on an unseen "
                  "day: neighbouring conditions from the same day leak information under random splits. For S1, heavily "
                  "regularised linear models (OLS/PLS/Ridge) generalise best across days, while trees and kNN overfit. For S2, "
                  "several models are worse than the mean on unseen days.")
image(s, "F4_model_zoo.png", 0.5, 1.55, w=8.4, h=5.35)
bullets(s, 9.15, 1.8, 3.65, 5.0, [
    [("Random CV: ", {"bold": True}), ("GP, SVR and the ensemble lead (MSE ≈ 0.008-0.011)", {})],
    [("Unseen day: ", {"bold": True}), ("error roughly doubles", {})],
    [("Simple wins across days: ", {"bold": True}), (f"PLS / Ridge reach {M('S1', 'PLS', 'day').MSE:.4f} on S1", {})],
    [("Trees & kNN ", {"bold": True}), ("overfit within-day structure", {})],
    [("Lesson: ", {"bold": True, "color": ORANGE}), ("report the protocol that matches deployment", {})],
], size=15, gap=12)

# ================================================================ 8. unseen day
s = content_slide("Predicting a session the model has never seen", "Results · generalisation",
                  "Each day is predicted by a GP trained only on the other four days; the shaded band is the GP 95% interval. "
                  "S1 tracks closely (r = 0.88, R² ≈ 0.79 for the best linear model). S2 misses the drop in level on days 3-5: "
                  "its sensors do not see that change, and the wider band shows the GP knows it is less sure. For CMA-ES, "
                  "ranking within a day is what matters; Spearman rho is about 0.6 for S1.")
image(s, "F5_unseen_day_timeline.png", 0.5, 1.5, w=8.6, h=5.4)
s1d, s2d = M("S1", "PLS", "day"), M("S2", "Elastic Net", "day")
stat(s, 9.35, 1.6, 3.4, "≈ 12 %", "error on unseen days, Subject 1 (R² 0.78)", BLUE, big_size=36)
stat(s, 9.35, 3.35, 3.4, f"ρ ≈ {s1d.within_day_spearman:.2f}", "within-day rank correlation (S1), "
     "what the optimiser needs", BLUE, big_size=34)
stat(s, 9.35, 5.1, 3.4, "≈ 16 %", "Subject 2: day-level shifts are invisible to the sensors", ORANGE, big_size=36)

# ================================================================ 9. features
s = content_slide("EMG carries the signal, and a few muscles do most of the work", "Results · interpretability",
                  "Permutation importance on held-out folds, plus a group-wise permutation. EMG > Step > Control for S1, and "
                  "EMG > Control > Step for S2, matching the paper's ablation. LASSO stability selection over 200 bootstraps "
                  "keeps R4 and R7 (S1) and R2 (S2) 93-100% of the time, so a reduced sensor set is realistic. Dropping the "
                  "control parameters barely hurts, which is good for portability to other devices.")
image(s, "F8_feature_importance.png", 0.5, 1.5, w=8.3, h=4.3)
abl = [("Feature set", "S1 GP", "S2 GP")]
for fs in ["All features", "Step + EMG", "EMG only", "Step only"]:
    g = lambda sub: e1[(e1.subject == sub) & (e1.model == "Gaussian Process") & (e1.feature_set == fs)].MSE.item()
    abl.append((fs, f"{g('S1'):.4f}", f"{g('S2'):.4f}"))
tb = s.shapes.add_table(len(abl), 3, Inches(9.1), Inches(1.65), Inches(3.65), Inches(2.3)).table
tb.columns[0].width = Inches(1.65)
for r, row in enumerate(abl):
    for c, v in enumerate(row):
        cell = tb.cell(r, c)
        cell.text = v
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.size, p.runs[0].font.name = Pt(14), BODY
        p.runs[0].font.bold = r == 0
        p.runs[0].font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (TINT if r % 2 else WHITE)
text(s, 9.1, 4.05, 3.65, 0.4, "Feature-set ablation (CV MSE)", size=11, color=INK2, italic=True)
bullets(s, 0.6, 6.0, 12.2, 0.9, [[("EMG is the strongest group ", {"bold": True}),
                                  ("in grouped importance for both subjects; dropping the controls costs little, which helps "
                                   "transfer to other exo hardware", {})]],
        size=15)

# ================================================================ 10. dimred + cross subject
s = content_slide("Why PCA failed, and why one model cannot fit everyone", "Results · deeper analysis",
                  "Left: PCA ignores the target. Supervised PLS with one component matches PCA with 8 on S2, which explains "
                  "the paper's poor PCA result. Right: train on one person, test on the other: worse than predicting the mean. "
                  "Per-subject z-scoring cuts the error 3x but is still not useful, so models must be subject-specific "
                  "(or need many more subjects plus anthropometrics).")
image(s, "F9_dimensionality_reduction.png", 0.5, 1.55, w=6.2, h=3.4)
image(s, "F12_cross_subject.png", 6.9, 1.55, w=5.9, h=3.4)
box(s, 0.6, 5.2, 5.95, 1.6)
text(s, 0.85, 5.33, 5.5, 1.4, [[("Supervised > unsupervised", {"bold": True, "color": NAVY, "size": 17})],
                               [("1 PLS component (0.0138) = 8 PCA components on S2. Forward selection is greedy "
                                 "and unstable.", {"size": 14, "color": INK2})]], spacing=4)
box(s, 6.9, 5.2, 5.85, 1.6)
text(s, 7.15, 5.33, 5.4, 1.4, [[("People are different", {"bold": True, "color": NAVY, "size": 17})],
                               [("S1 → S2: MSE 0.152 vs 0.036 for the mean. Per-subject z-scoring → 0.046. "
                                 "Personal models are needed.", {"size": 14, "color": INK2})]], spacing=4)

# ================================================================ 11. uncertainty + data
s = content_slide("The model knows when it is unsure, and more data would help", "Results · reliability",
                  "Calibration: the GP's nominal 95% interval covers 94% of unseen-day points for S1, close to perfect. For S2 "
                  "it covers 77%, so it is over-confident on the harder subject, but the intervals are visibly wider there. "
                  "Learning curves are still falling for S1 at 144 samples: more sessions should keep improving accuracy.")
image(s, "F11_calibration.png", 0.6, 1.5, w=4.4, h=3.75)
image(s, "F10_learning_curve.png", 5.2, 1.5, w=7.6, h=3.75)
stat(s, 0.6, 5.45, 4.4, "94 %", "of unseen-day points inside the S1 95 % interval", AQUA, big_size=32)
stat(s, 5.2, 5.45, 7.55, "0.0155 → 0.0088", "S1 GP test MSE from 28 to 144 training samples, still decreasing",
     AQUA, big_size=32)

# ================================================================ 11b. improvements
s = content_slide("Fixing the weaknesses, without collecting new data", "Improvements",
                  "Calibration: on a new day, measure k random conditions with the mask and shift that day's predictions "
                  "by a shrunken offset w * mean residual, with w = k / (k + sigma^2 / tau^2). Both variances are estimated "
                  "from the training days only (empirical Bayes). S2 trusts calibration (w about 0.87); S1, whose days "
                  "are already well predicted, trusts it less. 3 of 36 readings: S2 GP 0.046 -> 0.022. Calibrating on the "
                  "first conditions of a session is biased (cost runs high early on) and hurt S1 at k >= 5; shrinkage limits the damage. Sensors: "
                  "electrode choice is made inside each training fold, so there is no leakage. Cross-subject recalibration "
                  "did NOT help: a new person needs their own data.")
image(s, "F14_day_calibration.png", 0.5, 1.5, w=6.25, h=3.2)
image(s, "F15_sensor_reduction.png", 6.75, 1.5, w=6.1, h=3.2)
stat(s, 0.6, 4.95, 3.9, "−51 %", "S2 unseen-day error with just 3 of 36 mask spot-checks (GP 0.046 → 0.022)",
     ORANGE, big_size=34)
stat(s, 4.72, 4.95, 3.9, "w = k/(k+σ²/τ²)", "empirical-Bayes shrinkage: trusts calibration only as much as the "
     "training days justify", NAVY, big_size=24)
stat(s, 8.85, 4.95, 3.9, "4 of 16", "EMG electrodes beat all 16 for S2 (unseen-day Ridge 0.036 → 0.024)",
     AQUA, big_size=34)

# ================================================================ 12. demo
s = content_slide("Live demo: the code, running", "Demo",
                  "Open a terminal in the project folder (cd into ML proj). Run the three commands in order. (1) Likhith: the paper "
                  "reproduction prints linear-regression MSE 0.0107 and R2 0.83 for Subject 1. (2) Mukil: predict.py trains a "
                  "Gaussian Process on 4 days and predicts each unseen day; errors of about 11-17 %. (3) pytest shows all 10 "
                  "tests pass. Total under a minute.")
cmds = [("Paper reproduction", "python paper_reproduction/lin_reg_test.py",
         ["Linear Regression 5-Fold CV", "Mean MSE: 0.010700", "Mean R² : 0.8349"]),
        ("Predict unseen days", "python scripts/predict.py --demo --subject S1",
         ["day      MSE   %err      r", "  1   0.0286  16.9%   0.82", "  2   0.0191  13.8%   0.49",
          "  3   0.0132  11.5%   0.64", "  4   0.0119  10.9%   0.61", "  5   0.0149  12.2%   0.67"]),
        ("Automated tests", "python -m pytest -q", ["..........", "10 passed"])]
for i, (title, cmd, out) in enumerate(cmds):
    x = 0.6 + i * 4.1
    chip(s, x, 1.65, i + 1)
    text(s, x + 0.65, 1.7, 3.2, 0.45, title, size=17, bold=True, color=NAVY)
    box(s, x, 2.35, 3.85, 3.9, fill=NAVY, radius=0.05)
    text(s, x + 0.2, 2.5, 3.5, 0.9, "$ " + cmd, size=12, color=ICE, font="Courier New")
    text(s, x + 0.2, 3.45, 3.5, 2.7, [[(l, {})] for l in out], size=12, color=WHITE, font="Courier New")
text(s, 0.6, 6.45, 12.1, 0.4, "Run from the project folder:  cd \"ML proj\"   ·   total time under a minute",
     size=13, color=INK2, italic=True)

# ================================================================ 13. conclusions
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
_n[0] += 1
text(s, 0.6, 0.55, 12, 0.4, "CONCLUSIONS & FUTURE WORK", size=13, color=ICE, bold=True)
text(s, 0.6, 0.9, 12, 0.9, "Wearables can replace the mask, one person at a time",
     size=32, color=WHITE, font=HEAD, bold=True)
concl = [("9-10 %", "error within a session from EMG + gait (mean-baseline: 19-26 %); beats the original study"),
         ("2×", "error on a new day: random-split CV is optimistic, so deploy-style validation matters"),
         ("EMG", "is the key signal; 4 of 16 electrodes beat all 16 for S2, so cheaper hardware is possible"),
         ("1 : 1", "models must be subject-specific; cross-person transfer is worse than the mean")]
for i, (a, b) in enumerate(concl):
    x, y = 0.6 + (i % 2) * 6.2, 2.15 + (i // 2) * 1.55
    box(s, x, y, 5.95, 1.35, fill=NAVY2)
    text(s, x + 0.25, y + 0.2, 1.6, 0.9, a, size=30, color=ORANGE, font=HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.95, y + 0.17, 3.8, 1.05, b, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 5.45, 12, 0.4, "Next steps", size=17, color=WHITE, bold=True)
fut = ["More subjects + anthropometrics + IMU kinematics",
       "Test spot-check calibration live in a real HILO session",
       "Use the surrogate as a prior inside CMA-ES / Bayesian optimisation"]
for i, f in enumerate(fut):
    x = 0.6 + i * 4.1
    box(s, x, 5.9, 3.85, 0.9, fill=BLUE)
    text(s, x + 0.2, 5.9, 3.45, 0.9, f, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
s.notes_slide.notes_text_frame.text = (
    "Recap the four numbers. Future work: the biggest lever is more people. A daily calibration walk could fix the "
    "day-level offset seen in Subject 2. Use the prediction as a prior to speed up the optimiser rather than as a "
    "complete replacement for the mask.")

# ================================================================ 14. thank you / contributions
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.6, 1.2, 12, 1.0, "Thank you, questions?", size=44, color=WHITE, font=HEAD, bold=True)
text(s, 0.6, 2.35, 12, 0.5, "Code, data, figures and the 2-page write-up are in the private GitHub repository.",
     size=17, color=ICE)
for i, (who, what) in enumerate([("Mukil Skanda A · PES1UG24CS281", "Extended study: Bayesian-regularised NN, 13-model zoo, nested & "
                                                 "leave-one-day-out validation, calibration, sensor reduction, write-up"),
                                  ("Likhith M · PES1UG24CS282", "Faithful MATLAB → Python port of the paper (paper_reproduction/): "
                                                 "NN vs linear-regression baseline, 4 feature sets, forward selection, GitHub repo")]):
    x = 0.6 + i * 6.2
    box(s, x, 3.6, 5.95, 2.2, fill=NAVY2)
    text(s, x + 0.3, 3.8, 5.4, 0.5, who, size=20, color=ORANGE, bold=True)
    text(s, x + 0.3, 4.4, 5.4, 1.3, what, size=15, color=WHITE)
text(s, 0.6, 6.3, 12, 0.6, "Ref: Krimsky & Ng, CS229 2018 · Zhang et al., Science 2017 · MacKay 1992 · "
     "Rasmussen & Williams 2006", size=12, color=ICE)
s.notes_slide.notes_text_frame.text = "Mukil: extended study, calibration, sensor reduction. Likhith: faithful MATLAB-to-Python port and GitHub repo."

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
